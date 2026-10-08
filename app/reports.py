"""Reports built only from a run's stored results and snapshot, never from the live tables."""
from __future__ import annotations

import io
from collections import defaultdict
from dataclasses import dataclass, field

from sqlalchemy import select
from sqlalchemy.orm import Session

from . import models as m
from . import snapshot


@dataclass
class PatternSummary:
    rp: m.RunPattern
    logs: int = 0
    log_volume: float = 0.0
    dry_volume: float = 0.0
    wet_volume: float = 0.0
    value: float = 0.0
    sawdust: float = 0.0
    chips: float = 0.0
    boards: int = 0
    length_sum: float = 0.0
    mix: dict = field(default_factory=dict)

    def _per_log_volume(self, v: float) -> float:
        return v / self.log_volume if self.log_volume else 0.0

    @property
    def dry_recovery(self) -> float:
        return self._per_log_volume(self.dry_volume)

    @property
    def wet_recovery(self) -> float:
        return self._per_log_volume(self.wet_volume)

    @property
    def gross_value(self) -> float:
        """R of boards per m3 of log."""
        return self._per_log_volume(self.value)

    @property
    def residue_value(self) -> float:
        chips = self.chips * (1 - self.rp.pct_fines / 100.0) * self.rp.chip_price
        return self._per_log_volume(chips + self.sawdust * self.rp.sawdust_price)

    @property
    def nett_value(self) -> float:
        return self.gross_value + self.residue_value - self.rp.log_price

    @property
    def average_length(self) -> float:
        return self.length_sum / self.boards if self.boards else 0.0

    @property
    def boards_per_log(self) -> float:
        return self.boards / self.logs if self.logs else 0.0


def run_patterns(s: Session, run_id: int) -> list[PatternSummary]:
    rps = s.scalars(select(m.RunPattern).where(m.RunPattern.run_id == run_id).order_by(m.RunPattern.seq)).all()
    out = []
    for rp in rps:
        ps = PatternSummary(rp)
        for lr in s.scalars(select(m.RunLogResult).where(m.RunLogResult.run_pattern_id == rp.id)):
            ps.logs += 1
            ps.log_volume += lr.log_volume
            ps.dry_volume += lr.dry_volume
            ps.wet_volume += lr.wet_volume
            ps.value += lr.value
            ps.sawdust += lr.sawdust_volume
            ps.chips += lr.chip_volume
        for b in s.scalars(select(m.RunBoardResult).where(m.RunBoardResult.run_pattern_id == rp.id)):
            ps.boards += 1
            ps.length_sum += b.length_m
        out.append(ps)
    return out


def combined(patterns: list[PatternSummary]) -> PatternSummary:
    total = PatternSummary(m.RunPattern(line_name="", class_no=0, pattern_no=0, primary="All patterns", secondary="",
                                        log_price=0.0, chip_price=0.0, sawdust_price=0.0, pct_fines=0.0))
    for p in patterns:
        for f in ("logs", "log_volume", "dry_volume", "wet_volume", "value", "sawdust", "chips", "boards", "length_sum"):
            setattr(total, f, getattr(total, f) + getattr(p, f))
    # weighted price terms so the combined nett value is the volume-weighted nett value
    lv = total.log_volume or 1.0
    nett = sum(p.nett_value * p.log_volume for p in patterns) / lv
    total.rp.log_price = total.gross_value - nett
    return total


# ------------------------------------------------------------------ board report

LENGTH_MODES = {"none": "No lengths", "classified": "Length classes", "detailed": "Every length"}


def _length_label(length_m: float, classes: list[dict]) -> str:
    for c in classes:
        if c["min_m"] - 1e-6 <= length_m <= c["max_m"] + 1e-6:
            return f"{c['name']} ({c['min_m']:.1f}-{c['max_m']:.1f} m)"
    return "Other"


def board_report(s: Session, run: m.Run, pattern_id: int | None, lengths: str) -> list[dict]:
    """Rows of thickness x width (x length), pieces, volumes, value and shares."""
    snap = snapshot.loads(run.snapshot)
    classes = [{"name": c.name, "min_m": c.min_m, "max_m": c.max_m} for c in snap.products.length_classes]
    q = select(m.RunPattern.id).where(m.RunPattern.run_id == run.id)
    if pattern_id:
        q = q.where(m.RunPattern.id == pattern_id)
    ids = list(s.scalars(q))
    log_volume = sum(lr.log_volume for lr in s.scalars(select(m.RunLogResult).where(m.RunLogResult.run_pattern_id.in_(ids))))
    groups: dict[tuple, dict] = defaultdict(lambda: {"pieces": 0, "dry_volume": 0.0, "wet_volume": 0.0, "value": 0.0})
    graded = snap.products.grades_in_use
    for b in s.scalars(select(m.RunBoardResult).where(m.RunBoardResult.run_pattern_id.in_(ids))):
        if lengths == "detailed":
            key = (b.thickness, b.width, round(b.length_m, 2))
        elif lengths == "classified":
            key = (b.thickness, b.width, _length_label(b.length_m, classes))
        else:
            key = (b.thickness, b.width, "")
        key += (b.grade if graded else "",)
        g = groups[key]
        g["pieces"] += 1
        g["dry_volume"] += b.dry_volume
        g["wet_volume"] += b.wet_volume
        g["value"] += b.value
    sawn = sum(g["dry_volume"] for g in groups.values()) or 1.0
    rows = []
    for (t, w, ln, grade), g in sorted(groups.items(), key=lambda kv: (kv[0][0], kv[0][1], str(kv[0][2]), kv[0][3])):
        rows.append({"thickness": t, "width": w, "length": ln, "grade": grade, **g,
                     "share_sawn": g["dry_volume"] / sawn, "share_log": g["dry_volume"] / log_volume if log_volume else 0.0})
    return rows


def compare(a: list[PatternSummary], b: list[PatternSummary]) -> list[dict]:
    """Patterns of two runs side by side, matched on class and pattern text."""
    key = lambda p: (p.rp.class_no, p.rp.primary, p.rp.secondary)
    other = {key(p): p for p in b}
    rows = []
    for p in a:
        q = other.pop(key(p), None)
        rows.append({"a": p, "b": q})
    rows += [{"a": None, "b": q} for q in other.values()]
    return rows


# ------------------------------------------------------------------ summary

def volume_balance(p: PatternSummary) -> list[tuple[str, float, float]]:
    lv = p.log_volume or 1.0
    shrink = p.wet_volume - p.dry_volume
    rows = [("Sawn timber (dry)", p.dry_volume), ("Shrinkage", shrink), ("Chips", p.chips), ("Sawdust", p.sawdust)]
    rows.append(("Total", sum(v for _, v in rows)))
    return [("Logs", p.log_volume, 1.0)] + [(n, v, v / lv) for n, v in rows]


# ------------------------------------------------------------------ Excel

def excel(s: Session, run: m.Run) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    pats = run_patterns(s, run.id)
    wb = Workbook()
    bold = Font(bold=True)

    def sheet(ws, header, rows, formats=None):
        ws.append(header)
        for c in ws[1]:
            c.font = bold
        for r in rows:
            ws.append(list(r))
        for i, h in enumerate(header, 1):
            ws.column_dimensions[get_column_letter(i)].width = max(10, min(40, len(str(h)) + 2))
            fmt = (formats or {}).get(h)
            if fmt:
                for row in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                    row[0].number_format = fmt
        ws.freeze_panes = "A2"

    ws = wb.active
    ws.title = "One-liner"
    sheet(ws, ["Line", "Class", "Diameter (cm)", "Pattern", "Primary", "Secondary", "Logs", "Dry recovery",
               "Wet recovery", "Gross value (R/m³ log)", "Nett value (R/m³ log)", "Boards", "Boards per log",
               "Average length (m)", "Note"],
          [(p.rp.line_name, p.rp.class_no, p.rp.diameter_range, p.rp.pattern_no, p.rp.primary, p.rp.secondary, p.logs,
            p.dry_recovery, p.wet_recovery, p.gross_value, p.nett_value, p.boards, p.boards_per_log, p.average_length,
            p.rp.error) for p in pats],
          {"Dry recovery": "0.0%", "Wet recovery": "0.0%", "Gross value (R/m³ log)": "#,##0.00",
           "Nett value (R/m³ log)": "#,##0.00", "Boards per log": "0.00", "Average length (m)": "0.00"})

    ws = wb.create_sheet("Boards by pattern")
    rows = []
    for p in pats:
        for r in board_report(s, run, p.rp.id, "detailed"):
            rows.append((p.rp.class_no, p.rp.primary, p.rp.secondary, r["thickness"], r["width"], r["length"], r["grade"],
                         r["pieces"], r["dry_volume"], r["wet_volume"], r["value"], r["share_sawn"], r["share_log"]))
    sheet(ws, ["Class", "Primary", "Secondary", "Thickness (mm)", "Width (mm)", "Length (m)", "Grade", "Pieces", "Dry m³",
               "Wet m³", "Value (R)", "Share of sawn", "Share of log"], rows,
          {"Dry m³": "0.0000", "Wet m³": "0.0000", "Value (R)": "#,##0.00", "Share of sawn": "0.0%", "Share of log": "0.0%"})

    ws = wb.create_sheet("Boards combined")
    sheet(ws, ["Thickness (mm)", "Width (mm)", "Grade", "Pieces", "Dry m³", "Wet m³", "Value (R)", "Share of sawn",
               "Share of log"],
          [(r["thickness"], r["width"], r["grade"], r["pieces"], r["dry_volume"], r["wet_volume"], r["value"], r["share_sawn"],
            r["share_log"]) for r in board_report(s, run, None, "none")],
          {"Dry m³": "0.0000", "Wet m³": "0.0000", "Value (R)": "#,##0.00", "Share of sawn": "0.0%", "Share of log": "0.0%"})

    ws = wb.create_sheet("Summary")
    rows = []
    for p in pats + ([combined(pats)] if len(pats) > 1 else []):
        for name, v, share in volume_balance(p):
            rows.append((p.rp.class_no or "", p.rp.primary, p.rp.secondary, name, v, share))
    sheet(ws, ["Class", "Primary", "Secondary", "Item", "Volume (m³)", "Share of log volume"], rows,
          {"Volume (m³)": "0.0000", "Share of log volume": "0.0%"})

    ws = wb.create_sheet("Per log")
    rows = []
    for p in pats:
        for lr in s.scalars(select(m.RunLogResult).where(m.RunLogResult.run_pattern_id == p.rp.id).order_by(m.RunLogResult.log_no)):
            rows.append((p.rp.class_no, p.rp.primary, p.rp.secondary, lr.log_no, lr.sed_cm, lr.length_m, lr.log_volume,
                         lr.dry_volume, lr.wet_volume, lr.sawdust_volume, lr.chip_volume, lr.value, lr.boards))
    sheet(ws, ["Class", "Primary", "Secondary", "Log", "SED (cm)", "Length (m)", "Log m³", "Dry board m³",
               "Wet board m³", "Sawdust m³", "Chips m³", "Value (R)", "Boards"], rows,
          {"Log m³": "0.000000", "Dry board m³": "0.000000", "Wet board m³": "0.000000", "Sawdust m³": "0.000000",
           "Chips m³": "0.000000", "Value (R)": "#,##0.00"})

    ws = wb.create_sheet("Inputs")
    snap = snapshot.loads(run.snapshot)
    st = snap.settings
    info = [("Run", run.name), ("Created", run.created_at.isoformat(sep=" ") if run.created_at else ""),
            ("Calculated by", "Simsaw 6 (imported)" if run.source == "simsaw" else "This app"),
            ("Logs in snapshot", len(snap.logs)), ("Disc separation (cm)", st.disc_separation_cm),
            ("Nominal diameter used", "yes" if st.use_nominal_diameter else "no"),
            ("Nominal length used", "yes" if st.use_nominal_length else "no"),
            ("Nominal taper (mm/m)", st.nominal_taper_mm_per_m if st.use_nominal_taper else "actual")]
    for ln in snap.lines:
        info.append((f"{ln.name}: kerfs primary / secondary / edger (mm)",
                     f"{ln.primary_kerf:g} / {ln.secondary_kerf:g} / {ln.edger_kerf:g}"))
    sheet(ws, ["Item", "Value"], info)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
