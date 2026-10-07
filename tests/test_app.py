"""Web app: database round trip, import, runs, reports and every page, without a browser."""
import io
import pathlib

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import models as m
from app import reports, runs, snapshot, store
from app.db import Database, migrate
from app.main import create_app
from engine.sawing import simulate_pattern
from importers import simsaw

FIXTURES = pathlib.Path(__file__).parent / "fixtures"


@pytest.fixture
def db(tmp_path):
    url = f"sqlite:///{tmp_path / 'test.db'}"
    migrate(url)
    return Database(url)


@pytest.fixture
def ngomi(db):
    with db.session() as s:
        ds = store.import_simsaw(s, FIXTURES / "ngomi_1", "Ngomi")
        s.commit()
        return ds.id


@pytest.fixture
def client(tmp_path):
    app = create_app(f"sqlite:///{tmp_path / 'web.db'}", run_in_thread=False)
    with app.state.db.session() as s:
        ds = store.import_simsaw(s, FIXTURES / "ngomi_1", "Ngomi")
        s.commit()
    c = TestClient(app)
    c.ds = ds.id
    c.db = app.state.db
    return c


# ------------------------------------------------------------------ import and the engine view of the database

def test_imported_dataset_gives_the_engine_exactly_what_the_mdb_does(db, ngomi):
    ref = simsaw.load_inputs(FIXTURES / "ngomi_1")
    with db.session() as s:
        e = store.engine_dataset(s, ngomi)
    assert e.logs == ref.logs
    assert e.lines == ref.lines
    assert e.settings == ref.settings
    assert e.products.thicknesses == ref.products.thicknesses and e.products.widths == ref.products.widths
    assert sorted(e.products.combinations, key=repr) == sorted(ref.products.combinations, key=repr)
    assert e.products.wane == ref.products.wane
    assert [(c.no, c.min_diameter_cm, c.max_diameter_cm, c.log_price) for c in e.log_classes] == \
           [(c.no, c.min_diameter_cm, c.max_diameter_cm, c.log_price) for c in ref.log_classes]
    assert [(p.log_class_no, p.primary, p.secondary) for p in e.patterns] == \
           [(p.log_class_no, p.primary, p.secondary) for p in ref.patterns]
    assert [len(e.logs_in_class(c.no)) for c in e.log_classes] == [36, 34, 30, 44, 48]


def test_import_keeps_generator_grades_and_simsaw_run(db, ngomi):
    with db.session() as s:
        gen = s.query(m.LogGenerator).filter_by(dataset_id=ngomi).one()
        assert (gen.no_of_logs, gen.seed, gen.min_diameter, gen.max_diameter, gen.taper_distr) == (200, 1, 17.0, 41.9, 2)
        assert s.query(m.GradeOutput).filter_by(dataset_id=ngomi).count() == 16
        assert all(len(c.grades) == 1 for c in s.query(m.LogClass).filter_by(dataset_id=ngomi))
        run = s.query(m.Run).filter_by(dataset_id=ngomi).one()
        assert run.source == "simsaw" and run.status == "done"
        pats = reports.run_patterns(s, run.id)
    # Simsaw's own one-liner, recomputed from its stored per-log and per-board results
    assert [round(100 * p.dry_recovery, 1) for p in pats] == [54.0, 52.3, 56.2]
    assert [round(100 * p.wet_recovery, 1) for p in pats] == [61.6, 59.8, 64.0]
    assert [p.boards for p in pats] == [248, 217, 265]
    assert [round(p.nett_value, 2) for p in pats] == [2040.04, 1971.66, 2129.20]


def test_snapshot_round_trips_every_input():
    ds = simsaw.load_run(FIXTURES / "ngomi_1", "Test1").dataset
    assert snapshot.loads(snapshot.dumps(ds)) == ds


def test_duplicate_copies_everything_and_delete_removes_it(db, ngomi):
    with db.session() as s:
        copy = store.duplicate_dataset(s, ngomi, "Copy")
        s.commit()
        a, b = store.engine_dataset(s, ngomi), store.engine_dataset(s, copy.id)
        assert (a.logs, a.products, a.lines, a.log_classes, a.settings) == (b.logs, b.products, b.lines, b.log_classes, b.settings)
        assert [(p.primary, p.secondary) for p in a.patterns] == [(p.primary, p.secondary) for p in b.patterns]
        run = s.query(m.Run).filter_by(dataset_id=copy.id).one()
        assert [p.boards for p in reports.run_patterns(s, run.id)] == [248, 217, 265]
        store.delete_dataset(s, copy.id)
        s.commit()
        assert s.query(m.Log).filter_by(dataset_id=copy.id).count() == 0
        assert s.query(m.RunBoardResult).count() == 730            # the original run is untouched


def test_new_dataset_uses_ngomi_sizes_and_owner_wane_defaults_with_placeholders_marked(db):
    with db.session() as s:
        ds = store.create_default_dataset(s, "New")
        s.commit()
        e = store.engine_dataset(s, ds.id)
        assert [t.dry for t in e.products.thicknesses] == [19, 25, 38, 50]
        assert sum(c.valid for c in e.products.combinations) == 13
        assert e.products.wane[(38, 114)].thickness_pct == 0 and e.products.wane[(50, 152)].width_pct == 0
        assert (e.products.wane[(19, 76)].thickness_pct, e.products.wane[(25, 152)].width_pct) == (10, 30)
        assert all(c.price_placeholder for c in s.query(m.Combination).filter_by(dataset_id=ds.id))
        assert all(c.log_price_placeholder for c in s.query(m.LogClass).filter_by(dataset_id=ds.id))
        assert s.query(m.ProductionLine).filter_by(dataset_id=ds.id).one().kerfs_placeholder
        assert e.logs == [] and e.patterns == []


# ------------------------------------------------------------------ batch runs and reports

def test_batch_run_matches_the_engine_and_simsaw(db, ngomi):
    rid = runs.create_run(db, ngomi, "Check")
    runs.execute(db, rid)
    ref = simsaw.load_inputs(FIXTURES / "ngomi_1")
    with db.session() as s:
        run = s.get(m.Run, rid)
        assert run.status == "done" and run.progress == run.total == 106
        pats = reports.run_patterns(s, rid)
        for p, pd in zip(pats, ref.patterns):
            direct = simulate_pattern(ref.logs_in_class(pd.log_class_no), pd.primary, pd.secondary, ref.products,
                                      ref.lines[0], ref.settings, ref.log_class(pd.log_class_no).log_price)
            assert p.dry_recovery == pytest.approx(direct.dry_recovery, abs=1e-12)
            assert p.nett_value == pytest.approx(direct.nett_value_recovery, abs=1e-9)
            assert p.boards == direct.board_count
        simsaw_dry = [54.0, 52.29, 56.23]
        assert all(abs(100 * p.dry_recovery - ref) <= 1.0 for p, ref in zip(pats, simsaw_dry))
        # every log balances
        ids = [p.rp.id for p in pats]
        for lr in s.query(m.RunLogResult).filter(m.RunLogResult.run_pattern_id.in_(ids)):
            assert lr.wet_volume + lr.sawdust_volume + lr.chip_volume == pytest.approx(lr.log_volume, abs=1e-9)


def test_run_uses_its_snapshot_not_later_edits(db, ngomi):
    rid = runs.create_run(db, ngomi, "Before price change")
    with db.session() as s:
        for c in s.query(m.Combination).filter_by(dataset_id=ngomi):
            c.price = 1.0                                        # changed after the run was created
        s.commit()
    runs.execute(db, rid)
    with db.session() as s:
        assert reports.run_patterns(s, rid)[0].gross_value == pytest.approx(2164.16, abs=0.01)


def test_cancelled_run_stops_and_keeps_partial_results(db, ngomi):
    rid = runs.create_run(db, ngomi, "Stop me")
    runs._cancel[rid] = ev = __import__("threading").Event()
    ev.set()
    runs.execute(db, rid)
    with db.session() as s:
        assert s.get(m.Run, rid).status == "cancelled"


def test_runs_interrupted_by_closing_the_app_are_marked_on_start(tmp_path):
    url = f"sqlite:///{tmp_path / 'stop.db'}"
    migrate(url)
    db = Database(url)
    with db.session() as s:
        ds = store.create_default_dataset(s, "D")
        s.add(m.Run(dataset_id=ds.id, name="Half done", status="running"))
        s.commit()
    app = create_app(url, run_in_thread=False)
    with app.state.db.session() as s:
        assert s.query(m.Run).one().status == "cancelled"


def test_pattern_that_cannot_be_sawn_is_reported_not_fatal(db, ngomi):
    with db.session() as s:
        p = s.query(m.SawPattern).filter_by(dataset_id=ngomi).first()
        p.secondary = "19 22 38"
        s.commit()
    rid = runs.create_run(db, ngomi, "Bad pattern")
    runs.execute(db, rid)
    with db.session() as s:
        run = s.get(m.Run, rid)
        pats = reports.run_patterns(s, rid)
        assert run.status == "done" and "could not be sawn" in run.message
        assert "22 mm" in pats[0].rp.error and pats[1].boards == 217


def test_board_report_and_excel(db, ngomi):
    with db.session() as s:
        run = s.query(m.Run).filter_by(dataset_id=ngomi).one()
        rows = reports.board_report(s, run, None, "none")
        assert sum(r["pieces"] for r in rows) == 730
        assert sum(r["share_sawn"] for r in rows) == pytest.approx(1.0)
        detailed = reports.board_report(s, run, None, "detailed")
        assert sum(r["pieces"] for r in detailed) == 730 and len(detailed) > len(rows)
        classified = reports.board_report(s, run, None, "classified")
        assert {r["length"] for r in classified} == {"All (0.9-6.6 m)"}
        wb = load_workbook(io.BytesIO(reports.excel(s, run)))
    assert wb.sheetnames == ["One-liner", "Boards by pattern", "Boards combined", "Summary", "Per log", "Inputs"]
    one = list(wb["One-liner"].iter_rows(values_only=True))
    assert one[1][4:6] == ("25/114/25", "2*19 3*38 3*19") and round(one[1][7], 4) == 0.54


# ------------------------------------------------------------------ web pages and API

PAGES = ["/", "/d/{ds}", "/d/{ds}/logs", "/d/{ds}/products", "/d/{ds}/products?tab=prices", "/d/{ds}/products?tab=wane",
         "/d/{ds}/products?tab=centre", "/d/{ds}/products?tab=grades", "/d/{ds}/products?tab=residues",
         "/d/{ds}/machines", "/d/{ds}/machines?tab=general", "/d/{ds}/machines?tab=secondary",
         "/d/{ds}/machines?tab=edging", "/d/{ds}/settings", "/d/{ds}/patterns", "/d/{ds}/runs", "/d/{ds}/reports",
         "/d/{ds}/reports?view=boards&lengths=detailed", "/d/{ds}/reports?view=boards&pattern=",
         "/d/{ds}/reports?view=summary"]


def test_every_page_renders(client):
    for p in PAGES:
        r = client.get(p.format(ds=client.ds))
        assert r.status_code == 200, p


def test_grid_round_trip_validation_and_new_sizes(client):
    url = f"/api/d/{client.ds}/grid/thicknesses"
    rows = client.get(url).json()["rows"]
    bad = client.post(url, json={"rows": rows + [{"dry": "abc", "wet": ""}]})
    assert bad.status_code == 400 and len(bad.json()["errors"]) == 2
    dup = client.post(url, json={"rows": rows + [{"dry": "19", "wet": "21"}]})
    assert dup.status_code == 400
    ok = client.post(url, json={"rows": rows + [{"dry": "32", "wet": "35"}]})
    assert ok.status_code == 200 and [r["dry"] for r in ok.json()["rows"]] == [19, 25, 32, 38, 50]
    combos = client.get(f"/api/d/{client.ds}/grid/combinations").json()["rows"]
    new = [c for c in combos if c["thickness"] == 32]
    assert len(new) == 4 and all(c["price_placeholder"] and c["price"] == 0 for c in new)
    wane = [w for w in client.get(f"/api/d/{client.ds}/grid/wane").json()["rows"] if w["thickness"] == 32]
    assert all((w["thickness_pct"], w["width_pct"], w["length_wane"]) == (10, 30, 100) for w in wane)  # thinner than 38 mm


def test_editing_a_placeholder_price_clears_the_flag(client):
    with client.db.session() as s:
        ds = store.create_default_dataset(s, "Fresh")
        s.commit()
        dsid = ds.id
    url = f"/api/d/{dsid}/grid/combinations"
    rows = client.get(url).json()["rows"]
    rows[0]["price"] = "5200"
    rows[0]["price_placeholder"] = False
    r = client.post(url, json={"rows": rows}).json()["rows"]
    assert r[0]["price"] == 5200 and not r[0]["price_placeholder"] and r[1]["price_placeholder"]


def test_logs_grid_shows_class_and_accepts_pasted_rows(client):
    url = f"/api/d/{client.ds}/grid/logs"
    data = client.get(url).json()
    assert {r["log_no"]: r["class_no"] for r in data["rows"]}[8] == 1
    rows = data["rows"] + [{"log_no": "201", "sed_cm": "23,5", "length_m": "3", "taper": "9", "sweep_mm": "0",
                            "ovality": "1", "defect_core_cm": "0", "log_grade_id": "All log grades"}]
    out = client.post(url, json={"rows": rows}).json()["rows"]
    assert out[-1]["sed_cm"] == 23.5 and out[-1]["class_no"] == 2


def test_diagram_api_matches_the_worked_example(client):
    with client.db.session() as s:
        line = s.query(m.ProductionLine).filter_by(dataset_id=client.ds).one().id
    d = client.get(f"/api/d/{client.ds}/diagram", params={"line_id": line, "log_no": 8, "primary": "25/114/25",
                                                          "secondary": "2*19 3*38 3*19"}).json()
    assert d["problems"] == []
    assert (d["cant"]["lo"], d["cant"]["hi"]) == (-60, 60)
    assert [63.0, 90.0] in [[b["left"], b["right"]] for b in d["boards"]]
    stack = [(b["bottom"], b["top"]) for b in d["boards"] if b["thickness"] == 38]
    assert stack == [(-76.5, -35.5), (-32.5, 8.5), (11.5, 52.5)]
    assert d["result"]["log_volume"] == pytest.approx(0.105927, abs=1e-6)
    assert d["blades"]["primary"] == [-93, -90, -63, -60, 60, 63, 90, 93]
    bad = client.get(f"/api/d/{client.ds}/diagram", params={"line_id": line, "log_no": 8, "primary": "25/114/25",
                                                            "secondary": "22"}).json()
    assert "22 mm" in bad["problems"][0] and "large_end" in bad


def test_class_simulation_and_pattern_save(client):
    with client.db.session() as s:
        line = s.query(m.ProductionLine).filter_by(dataset_id=client.ds).one().id
        cls = s.query(m.LogClass).filter_by(dataset_id=client.ds, no=1).one().id
    r = client.get(f"/api/d/{client.ds}/simulate_class", params={"line_id": line, "class_id": cls,
                                                                 "primary": "25/114/25", "secondary": "2*19 3*38 3*19"}).json()
    assert r["logs"] == 36 and r["boards"] == 248 and round(100 * r["dry_recovery"], 2) == 54.10
    new = client.post(f"/api/d/{client.ds}/patterns", json={"line_id": line, "class_id": cls,
                                                            "primary": "25/102/25", "secondary": "3*38"}).json()
    listing = client.get(f"/api/d/{client.ds}/patterns", params={"line_id": line, "class_id": cls}).json()
    assert [p["pattern_no"] for p in listing["patterns"]] == [1, 2, 3] and listing["patterns"][-1]["id"] == new["id"]
    assert len(listing["logs"]) == 36
    client.post(f"/api/d/{client.ds}/patterns/{new['id']}/delete")
    assert len(client.get(f"/api/d/{client.ds}/patterns", params={"line_id": line, "class_id": cls}).json()["patterns"]) == 2


def test_run_from_the_web_and_download_excel(client):
    r = client.post(f"/d/{client.ds}/runs", data={"name": "Web run"}, follow_redirects=False)
    assert r.status_code == 303
    with client.db.session() as s:
        run = s.query(m.Run).filter_by(name="Web run").one()
        assert run.status == "done"
    st = client.get(f"/api/runs/{run.id}").json()
    assert st["progress"] == st["total"] == 106
    x = client.get(f"/d/{client.ds}/reports/{run.id}/excel")
    assert x.status_code == 200 and x.content[:2] == b"PK"
    assert "54.1 %" in client.get(f"/d/{client.ds}/reports?run={run.id}").text


def test_log_generator_from_the_web_is_repeatable(client):
    form = {"no_of_logs": "50", "seed": "7", "fixed_seed": "on", "mode": "replace", "length_incr": "0.3",
            **{f"{a}_{k}": v for k, (lo, hi) in {"diameter": (18, 40), "length": (2.4, 3.0), "taper": (8, 11),
                                                 "sweep": (0, 15), "ovality": (0.95, 1.05), "defect_core": (0, 0)}.items()
               for a, v in (("min", lo), ("max", hi))},
            **{f"{k}_distr": "0" for k in ("diameter", "length", "taper", "sweep", "ovality", "defect_core")}}
    client.post(f"/d/{client.ds}/logs/generate", data=form)
    first = client.get(f"/api/d/{client.ds}/grid/logs").json()["rows"]
    client.post(f"/d/{client.ds}/logs/generate", data=form)
    second = client.get(f"/api/d/{client.ds}/grid/logs").json()["rows"]
    strip = lambda rows: [{k: v for k, v in r.items() if k != "id"} for r in rows]
    assert len(first) == 50 and strip(first) == strip(second)


def test_machine_form_saves_and_clears_kerf_placeholder(client):
    with client.db.session() as s:
        ln = s.query(m.ProductionLine).filter_by(dataset_id=client.ds).one()
        ln.kerfs_placeholder = True
        s.commit()
        lid = ln.id
    client.post(f"/d/{client.ds}/machines/{lid}", data={"tab": "primary", "primary_kerf": "3.4",
                                                         "primary_outside_kerf": "", "primary_outside_blades": "0"})
    with client.db.session() as s:
        ln = s.get(m.ProductionLine, lid)
        assert ln.primary_kerf == 3.4 and ln.primary_outside_kerf is None and not ln.kerfs_placeholder


def test_dataset_new_duplicate_delete_from_the_web(client):
    r = client.post("/datasets/new", data={"name": "Trial"}, follow_redirects=False)
    new_id = int(r.headers["location"].rsplit("/", 1)[1])
    r = client.post(f"/datasets/{new_id}/duplicate", data={"name": "Trial copy"}, follow_redirects=False)
    copy_id = int(r.headers["location"].rsplit("/", 1)[1])
    assert "Trial copy" in client.get("/").text
    client.post(f"/datasets/{copy_id}/delete")
    assert "Trial copy" not in client.get("/").text
    bad = client.post("/datasets/import", files={"file": ("notes.txt", b"hello")}, follow_redirects=True)
    assert "Choose a Simsaw .mdb file" in bad.text
